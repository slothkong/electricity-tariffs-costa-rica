import yaml
import pandas
import numpy

import os
import requests

from pandas import DataFrame
from numpy import array

from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell


WORKBOOK_URL = os.getenv("WORKBOOK_URL", "https://aresep-my.sharepoint.com/:x:/g/personal/multimedia_aresep_go_cr/ET6L4k-QyphAgLpEwSYeNegBbLvOGM7mF0n2vZxId_SGeQ?e=YGhFZr&download=1")
WORKBOOK_PATH =  os.getenv("WORKBOOK_PATH", "./kaggle/input/datasets/slothkong/tarifas-electricas-de-costa-rica/Cuadro E-8 Tarifas electricas final.xlsx")
CONFIG_PATH = os.getenv("CONFIG_PATH", "./kaggle/working/cfgs.yaml")
DATAFRAME_PATH = os.getenv("DATAFRANE_PATH", "./kaggle/working/electricity-tariffs-costa-rica.csv")
MONTH_MAPPING = {"enero": "01", "febrero": "02", "marzo": "03", "abril": "04", "mayo": "05", "junio": "06", "julio": "07", "agosto": "08", "septiembre": "09", "octubre": "10", "noviembre": "11", "diciembre": "12"}

def download_workbook(workbook_url: str, workbook_path: str) -> None:
    
    response = requests.get(workbook_url, stream=True)
    response.raise_for_status()

    with open(workbook_path, "wb") as file:
        for chunk in response.iter_content(chunk_size=8192):
            file.write(chunk)

def select_worksheet(worksheet_title: str, workbook: Workbook) -> Worksheet:

    def worksheet_lookup(workbook: Workbook)-> dict[str, int]:
        worksheet_lookup_map = {}

        for idx, worksheet in enumerate(workbook.worksheets):
            worksheet_lookup_map.update({worksheet.title: idx})

        return worksheet_lookup_map
    
    idx = worksheet_lookup(workbook).get(worksheet_title)

    if not idx:
        raise ValueError(f"Desired worksheet title '{worksheet_title}' not in workbook.")

    return workbook.worksheets[idx]

def pivot_table(cell_range: tuple[Cell,])-> DataFrame:

    column_names = ["bloque_de_tarifa"] + [month for month in MONTH_MAPPING.keys()]
    column_types_mapping = dict({column_names[0]: "str"}) | dict(zip(column_names[1:], ["float64"] * len(column_names[1:])))

    for row_idx in range(len(cell_range)):
        values = []
        for row in cell_range[row_idx]:
    
            values.append([row.value])
    
        record = DataFrame.from_records(
            data=numpy.array(values).T,
            columns=column_names
        )

        if record.empty:
            raise ValueError("Unable to extract valid data from user-defined cell range!")

        record = record.astype(column_types_mapping)
        
        if row_idx == 0:
            dataframe = record
        else:
            dataframe = pandas.concat([dataframe, record], ignore_index=True)
    
    dataframe = pandas.melt(dataframe,
                            id_vars=[column_names[0]],
                            value_vars=column_names[1:],
                            var_name="mes",
                            value_name="colones_por_kwh")    

    return dataframe

def convert_workbook_to_dataframe(workbook_path: str, config_path: str) -> DataFrame:

    column_names = ["annio_mes", "distribuidor", "tipo_de_tarifa", "bloque_de_tarifa", "colones_por_kwh"]
    is_initial_loop = True

    with open(config_path, "r") as fp:
        cfgs = yaml.safe_load(fp)

    workbook = load_workbook(workbook_path, data_only=True)
    for worksheet_title, worksheets_cfg in cfgs["worksheets"].items():

        worksheet = select_worksheet(worksheet_title, workbook)
        for tariff_type, worksheet_cfg in worksheets_cfg["ranges"].items():
            
            cell_range = worksheet_cfg["range"]
            print(f"Processing range '{cell_range}' from worksheet '{worksheet_title}'...")

            year = worksheet_title.split(" ")[1]
            distributor_acronym = worksheet_title.split(" ")[0]

            tmp_dataframe = pivot_table(worksheet[cell_range])
            tmp_dataframe["annio"] = year
            tmp_dataframe["tipo_de_tarifa"] = tariff_type
            tmp_dataframe["distribuidor"] = distributor_acronym
            tmp_dataframe["annio_mes"] = tmp_dataframe["mes"].apply(lambda x: f"{year}-{MONTH_MAPPING.get(x)}")
            tmp_dataframe = tmp_dataframe[column_names]

            if is_initial_loop:
                dataframe = tmp_dataframe
                is_initial_loop = False
            else:
                dataframe = pandas.concat([dataframe, tmp_dataframe], ignore_index=True)


    dataframe.sort_values(by=column_names[0:-1], inplace=True)
    return dataframe

def main() -> None:
    
    download_workbook(WORKBOOK_URL, WORKBOOK_PATH)
    print(f"Successfully downloaded input workbook from '{WORKBOOK_URL}'")
    
    dataframe = convert_workbook_to_dataframe(WORKBOOK_PATH, CONFIG_PATH)
    
    dataframe.to_csv(DATAFRAME_PATH, index=False)
    print(f"Successfully wrote dataframe to '{DATAFRAME_PATH}'")


if __name__ == "__main__":
    main()
